#!/usr/bin/env python3
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

W = Inches(13.33); H = Inches(7.5)
F = "/Users/lorenajakic/Desktop/LaTeX-template-for-final-thesis-at-UniZG-FER/Quickstart Examples/Figures/"

INDIGO=RGBColor(0x4F,0x46,0xE5); PURPLE=RGBColor(0x7C,0x3A,0xED)
ORANGE=RGBColor(0xEA,0x76,0x0C); GREEN=RGBColor(0x16,0xA3,0x4A)
CORAL=RGBColor(0xF4,0x3F,0x5E); TEAL=RGBColor(0x0E,0xA5,0xE9)
AMBER=RGBColor(0xF5,0x9E,0x0B); DARK=RGBColor(0x1E,0x29,0x3B)
GRAY=RGBColor(0x64,0x74,0x8B); LGRAY=RGBColor(0xF1,0xF5,0xF9)
WHITE=RGBColor(0xFF,0xFF,0xFF)
LPUR=RGBColor(0xED,0xE9,0xFE); LORANGE=RGBColor(0xFF,0xED,0xD5)
LGREEN=RGBColor(0xDC,0xFC,0xE7); LCORAL=RGBColor(0xFF,0xE4,0xE6)
LTEAL=RGBColor(0xE0,0xF2,0xFE); LIND=RGBColor(0xE0,0xE7,0xFF)

FONT = "Arial"

def rect(sl,x,y,w,h,fill=None,line=None,lw=None,rnd=False):
    sh=sl.shapes.add_shape(5 if rnd else 1,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    else: sh.fill.background()
    if line: sh.line.color.rgb=line; sh.line.width=Pt(lw or 1.5)
    else: sh.line.fill.background()
    return sh

def oval(sl,x,y,w,h,fill=None,line=None,lw=2.0):
    sh=sl.shapes.add_shape(9,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    else: sh.fill.background()
    if line: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    return sh

def txt(sl,x,y,w,h,s,size=20,bold=False,color=None,align=PP_ALIGN.LEFT,italic=False):
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=s; r.font.size=Pt(size); r.font.bold=bold
    r.font.italic=italic; r.font.color.rgb=color or DARK
    r.font.name=FONT
    return tb

def lbl(sl,x,y,w,h,s,size=18,bold=False,color=None,bg=None,align=PP_ALIGN.CENTER,rnd=False):
    sh=rect(sl,x,y,w,h,fill=bg,rnd=rnd)
    tf=sh.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=s; r.font.size=Pt(size); r.font.bold=bold
    r.font.color.rgb=color or DARK; r.font.name=FONT
    return sh

def connector(sl,x1,y1,x2,y2,color=GRAY,w=1.5,dash=False):
    c=sl.shapes.add_connector(1,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(w)
    if dash: c.line.dash_style=4
    return c

def img(sl,path,x,y,w=None,h=None):
    kw={}
    if w: kw['width']=Inches(w)
    if h: kw['height']=Inches(h)
    return sl.shapes.add_picture(path,Inches(x),Inches(y),**kw)

def header(sl,title,sub=None,color=INDIGO):
    rect(sl,0,0,0.12,7.5,fill=color)
    txt(sl,0.3,0.15,12.6,0.8,title,size=34,bold=True,color=DARK)
    if sub: txt(sl,0.3,0.92,12.6,0.45,sub,size=18,color=GRAY,italic=True)
    rect(sl,0.3,1.35,12.75,0.04,fill=LGRAY)

# ── animation ────────────────────────────────────────────────
_c=[10]
def nid(): v=_c[0]; _c[0]+=1; return str(v)
def reset(): _c[0]=10

def _click_group(elements):
    outer=nid()
    inner_xml=""
    for i,(spid,atype) in enumerate(elements):
        node="clickEffect" if i==0 else "withEffect"
        a,b,c=nid(),nid(),nid()
        if atype=='fade':
            child=f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="500" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
            pre=f'presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}" dur="500"'
        else:
            child=f'<p:set><p:cBhvr><p:cTn id="{c}" dur="1" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            pre=f'presetID="1" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="{node}"'
        inner_xml+=f'<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{b}" {pre}><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{child}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
    return parse_xml(f'<p:par xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond evt="onClick" delay="0"><p:tn><p:tgtEl><p:sldTgt/></p:tgtEl></p:tn></p:cond></p:stCondLst><p:childTnLst>{inner_xml}</p:childTnLst></p:cTn></p:par>')

def animate(sl, groups):
    reset()
    r,s=nid(),nid()
    timing=parse_xml(f'<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="{r}" dur="indefinite" restart="whenNotActive" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="{s}" dur="indefinite" nodeType="mainSeq"><p:childTnLst/></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tn><p:tgtEl><p:sldTgt/></p:tgtEl></p:tn></p:cond></p:prevCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst/></p:timing>')
    cl=timing.find('.//' + qn('p:seq') + '/' + qn('p:cTn') + '/' + qn('p:childTnLst'))
    for g in groups:
        cl.append(_click_group([(sh.shape_id, at) for sh,at in g]))
    sl._element.append(timing)

prs=Presentation()
prs.slide_width=W; prs.slide_height=H
def sl(): return prs.slides.add_slide(prs.slide_layouts[6])

# ═══════════════════════════════════════════════════════════════
# SLIDE 1 — Naslov
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
rect(s,8.5,0,4.83,7.5,fill=RGBColor(0xF5,0xF3,0xFF))
rect(s,0,0,0.55,7.5,fill=INDIGO)
img(s,F+"nodes-1-large.png",8.3,0.5,w=4.8)
txt(s,0.8,0.7,7.4,3.0,"Aplikacija za planiranje\nputovanja primjenom\nLLM-agenta",size=40,bold=True,color=DARK)
txt(s,0.8,3.75,7.1,0.7,"Od kaosa grupnog dogovaranja\ndo sinkroniziranog agentskog plana",size=20,color=GRAY,italic=True)
rect(s,0.8,4.6,5.4,0.05,fill=INDIGO)
txt(s,0.8,4.8,5.5,0.5,"Lorena Jakić",size=22,bold=True,color=DARK)
txt(s,0.8,5.32,7.1,0.4,"Mentor: doc. dr. sc. Luka Humski",size=18,color=GRAY)
txt(s,0.8,5.75,7.2,0.4,"FER Zagreb — lipanj 2026.",size=18,color=GRAY)
lbl(s,0.8,6.4,2.5,0.6,"GoTogether",size=20,bold=True,color=WHITE,bg=INDIGO,rnd=True)

# ═══════════════════════════════════════════════════════════════
# SLIDE 2 — Pregled (Agenda)
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
rect(s,0,0,0.12,7.5,fill=INDIGO)
txt(s,0.3,0.15,12.6,0.8,"Pregled prezentacije",size=34,bold=True,color=DARK)
rect(s,0.3,0.95,12.75,0.04,fill=LGRAY)

TOPICS=[
    ("1","Motivacija","Zašto ova aplikacija?",CORAL,LCORAL),
    ("2","Rješenje","Što je GoTogether?",TEAL,LTEAL),
    ("3","Agentski sustav","Kako AI agent radi?",PURPLE,LPUR),
    ("4","Demo","Aplikacija u praksi",ORANGE,LORANGE),
    ("5","Evaluacija","Rezultati i zaključak",GREEN,LGREEN),
]
anim=[]
for i,(num,ti,su,col,bg) in enumerate(TOPICS):
    cx=1.4+i*2.42
    o=oval(s,cx-0.65,2.3,1.3,1.3,fill=bg,line=col,lw=2.5)
    n=oval(s,cx-0.38,2.57,0.76,0.76,fill=col)
    n.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
    run=n.text_frame.paragraphs[0].add_run()
    run.text=num; run.font.size=Pt(22); run.font.bold=True
    run.font.color.rgb=WHITE; run.font.name=FONT
    t=lbl(s,cx-1.15,3.8,2.3,0.6,ti,size=18,bold=True,color=col)
    su_=lbl(s,cx-1.15,4.45,2.3,0.65,su,size=16,color=GRAY)
    anim.append([(o,'fade'),(n,'appear'),(t,'appear'),(su_,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 3 — Motivacija: Problem
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Motivacija: kako planiramo putovanje s prijateljima?",
       "Zamislite: 6 prijatelja, Pariz, 5 dana...",color=CORAL)

items=[
    ("💬","Dogovaranje na 5 mjesta",
     "Facebook, Instagram, e-mail, Messenger...\nSvaki kanal drugačija odluka"),
    ("🎟️","Karte i rezervacije razbacane",
     "Hoteli u mailu, letovi u appu,\nvlakske karte u grupnom chatu"),
    ("⏰","Nitko ne provjerava izvedivost",
     "Plan izgleda dobro na papiru —\nali termini se preklapaju"),
    ("❌","Nema jednog mjesta za sve",
     "Prijedlozi se gube,\nodluke se ponavljaju"),
]
anim=[]
for i,(icon,ti,de) in enumerate(items):
    y=1.5+i*1.38
    ic=lbl(s,0.28,y,0.75,0.75,icon,size=24,bg=LCORAL,color=CORAL,rnd=True)
    t=txt(s,1.18,y+0.02,5.5,0.45,ti,size=20,bold=True,color=DARK)
    d=txt(s,1.18,y+0.5,5.5,0.78,de,size=17,color=GRAY)
    anim.append([(ic,'appear'),(t,'appear'),(d,'appear')])

q=rect(s,7.0,1.5,6.1,4.1,fill=RGBColor(0xFF,0xF7,0xF7),rnd=True)
q.line.color.rgb=CORAL; q.line.width=Pt(1.5)
tf=q.text_frame; tf.word_wrap=True
p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
r=p.add_run(); r.font.size=Pt(19); r.font.color.rgb=DARK; r.font.name=FONT
r.text=('💬  "Kada idemo?\nTko je pronašao hotel?\nJeste li vidjeli moju poruku o kartama?\nČekajte — koji plan smo prihvatili?"\n\n'
        '— Tipični grupni chat')

sol=lbl(s,7.0,5.72,6.1,1.6,
    "✅  Treba nam jedno mjesto\ngdje svi planiraju zajedno,\na AI koordinira i provjerava.",
    size=18,bold=False,color=WHITE,bg=CORAL,rnd=True)
anim.append([(q,'fade'),(sol,'fade')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 4 — Rješenje: GoTogether
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Rješenje: GoTogether",
       "Jedno mjesto — svi planiraju zajedno, AI asistent provjerava svaki korak",color=INDIGO)

img(s,F+"cijeli_trip.png",0.28,1.5,w=6.9)

features=[
    (INDIGO,LIND,"🗺️  Interaktivna karta","Dodaj mjesta klikom — vidljivo svima odmah"),
    (PURPLE,LPUR,"🤖  AI asistent","Generira plan, traži letove, smještaj, odgovara na pitanja"),
    (GREEN,LGREEN,"📂  Svi dokumenti","Karte, rezervacije i potvrde na jednom mjestu"),
    (TEAL,LTEAL,"👥  Suradnja u stvarnom vremenu","Svaka izmjena vidljiva svim sudionicima odmah"),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(features):
    y=1.5+i*1.48
    b=rect(s,7.4,y,5.72,1.35,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.0)
    t=txt(s,7.58,y+0.1,5.35,0.48,ti,size=18,bold=True,color=col)
    d=txt(s,7.58,y+0.62,5.35,0.65,de,size=16,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 5 — Arhitektura sustava
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Arhitektura sustava",
       "Dvije glavne komponente komuniciraju HTTP protokolom",color=TEAL)

img(s,F+"sustav.png",0.28,1.5,w=7.6)

boxes=[
    (8.1,1.5,INDIGO,LIND,"Web-aplikacija\n(Ruby on Rails 8)",
     "MVC · Devise · Turbo Streams · WebSocket"),
    (8.1,3.3,TEAL,LTEAL,"PostgreSQL",
     "Korisnici · putovanja · mjesta · poruke"),
    (8.1,5.05,PURPLE,LPUR,"Agentski sustav\n(Python / LangGraph)",
     "Višestupanjski AI agent · HTTP veza"),
    (10.75,5.05,CORAL,LCORAL,"Vanjski servisi",
     "Skyscanner · Booking · Airbnb · Tavily"),
]
anim=[]
for x,y,col,bg,ti,de in boxes:
    b=rect(s,x,y,2.38,1.68,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.2)
    t=txt(s,x+0.12,y+0.1,2.12,0.55,ti,size=15,bold=True,color=col)
    d=txt(s,x+0.12,y+0.68,2.12,0.85,de,size=13,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 6 — Zašto agent, ne samo LLM?
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Zašto nije dovoljno samo pitati LLM?",
       "Jedan poziv vs. višestupanjski agentski sustav",color=PURPLE)

lbl(s,0.22,1.45,6.2,0.6,"❌  Trivijalni pristup (Zero-shot)",
    size=18,bold=True,color=DARK,bg=LGRAY,rnd=True)
trivial=[
    "Jedan poziv → jedan odgovor",
    "Ne provjerava preklapanja termina",
    "Ignorira radna vremena lokacija",
    "Ne zna stvarne udaljenosti",
    "Plan izgleda ispravno — ali nije",
]
for i,t_ in enumerate(trivial):
    y=2.2+i*0.97
    ic=lbl(s,0.22,y,0.65,0.65,"✗",size=22,bold=True,color=CORAL,bg=LCORAL,rnd=True)
    txt(s,1.0,y+0.1,5.3,0.65,t_,size=18,color=DARK)

rect(s,6.58,1.42,0.06,5.8,fill=LGRAY)

lbl(s,6.75,1.45,6.4,0.6,"✅  Agentski sustav (LangGraph)",
    size=18,bold=True,color=PURPLE,bg=LPUR,rnd=True)
agentic=[
    "Više koraka → iterativno poboljšanje",
    "Deterministička provjera ograničenja",
    "Stvarne cijene letova i smještaja",
    "Petlja popravka dok plan nije valjan",
    "Bodovanje → bira se najbolji plan",
]
for i,t_ in enumerate(agentic):
    y=2.2+i*0.97
    ic=lbl(s,6.75,y,0.65,0.65,"✓",size=22,bold=True,color=GREEN,bg=LGREEN,rnd=True)
    txt(s,7.52,y+0.1,5.55,0.65,t_,size=18,color=DARK)

# ═══════════════════════════════════════════════════════════════
# SLIDE 7 — Graf agentskog sustava
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Agentski sustav: Kako agent odlučuje?",
       "Graf čvorova — svaki čvor = jedan korak obrade",color=PURPLE)

def nd_oval(s,cx,cy,w,h,label,fill,line,size=13,bold=True):
    sh=oval(s,cx-w/2,cy-h/2,w,h,fill=fill,line=line,lw=2.0)
    tf=sh.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=label; r.font.size=Pt(size)
    r.font.bold=bold; r.font.color.rgb=DARK; r.font.name=FONT
    return sh
def nd_rect(s,cx,cy,w,h,label,fill,line,size=13):
    sh=rect(s,cx-w/2,cy-h/2,w,h,fill=fill,line=line,lw=2.0,rnd=True)
    tf=sh.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
    r=p.add_run(); r.text=label; r.font.size=Pt(size)
    r.font.bold=True; r.font.color.rgb=DARK; r.font.name=FONT
    return sh

GX=4.3; GYS=1.85; GYR=2.85; GYB=4.3; GYP=5.55; GYE=6.75
RX=7.5; RS=1.5

sh_s =nd_rect(s,GX,GYS,1.8,0.6,"__početak__",RGBColor(0xBB,0xF7,0xD0),GREEN,size=12)
sh_r =nd_oval(s,GX,GYR,1.7,0.8,"usmjeravanje",LPUR,PURPLE,size=12)
sh_np=nd_oval(s,2.4,GYB,1.55,0.75,"novi plan",LPUR,PURPLE,size=11)
sh_ip=nd_oval(s,4.3,GYB,1.65,0.75,"izmjena\nplana",LPUR,PURPLE,size=10)
sh_pl=nd_oval(s,3.3,GYP,1.65,0.8,"izrada\nplana",RGBColor(0xFF,0xED,0xD5),ORANGE,size=11)
sh_rc=nd_oval(s,RX,GYB,1.5,0.72,"preporuke",LPUR,PURPLE,size=10)
sh_ac=nd_oval(s,RX+RS,GYB,1.5,0.72,"smještaj",LPUR,PURPLE,size=10)
sh_tr=nd_oval(s,RX+2*RS,GYB,1.5,0.72,"prijevoz",LPUR,PURPLE,size=10)
sh_gn=nd_oval(s,RX+3*RS,GYB,1.62,0.72,"opća\npitanja",LPUR,PURPLE,size=10)
sh_e =nd_rect(s,GX+1.5,GYE,1.8,0.6,"__kraj__",RGBColor(0xBB,0xF7,0xD0),GREEN,size=12)

def a(sl,x1,y1,x2,y2,d=False): connector(sl,x1,y1,x2,y2,color=DARK,w=1.5,dash=d)
a(s,GX,GYS+0.3,GX,GYR-0.4)
a(s,GX-0.4,GYR+0.4,2.4,GYB-0.38)
a(s,GX,GYR+0.4,4.1,GYB-0.38)
a(s,GX+0.35,GYR+0.38,RX-0.5,GYB-0.36)
a(s,GX+0.55,GYR+0.32,RX+RS-0.4,GYB-0.36)
a(s,GX+0.7,GYR+0.25,RX+2*RS-0.3,GYB-0.36)
a(s,GX+0.82,GYR+0.18,RX+3*RS-0.15,GYB-0.36)
a(s,2.4,GYB+0.38,3.0,GYP-0.4)
a(s,4.3,GYB+0.38,3.6,GYP-0.4)
a(s,3.3,GYP+0.4,GX+0.9,GYE-0.3,True)
for rx in [RX,RX+RS,RX+2*RS,RX+3*RS]:
    a(s,rx,GYB+0.38,GX+2.1,GYE-0.3,True)

callouts=[
    (sh_s, 9.8,1.42, LGREEN,GREEN,"__početak__",
     "Ulazna točka — svaki korisnički upit počinje ovdje"),
    (sh_r, 9.8,2.65, LPUR,PURPLE,"usmjeravanje",
     "LLM klasificira upit u jednu od 6 kategorija\n(brzi claude-haiku model)"),
    (sh_np,9.8,3.88, LORANGE,ORANGE,"novi plan / izmjena plana",
     "Paralelno generira 2 kandidata:\nkompaktnost vs. raznolikost aktivnosti"),
    (sh_pl,9.8,5.1,  LORANGE,ORANGE,"izrada plana (podgraf)",
     "Provjera → popravak → bodovanje\nBira se najkvalitetniji valjani plan"),
    (sh_rc,9.8,6.3,  LPUR,PURPLE,"preporuke / smještaj / prijevoz",
     "Pozivaju vanjske API-je u stvarnom vremenu\n(Skyscanner, Booking, Airbnb, Tavily)"),
]
anim=[]
for sh_n,bx,by,bg,col,ti,de in callouts:
    nx=sh_n.left/914400; ny=sh_n.top/914400
    nw=sh_n.width/914400; nh=sh_n.height/914400
    ring=oval(s,nx-0.12,ny-0.12,nw+0.24,nh+0.24,line=col,lw=3.0)
    ring.fill.background()
    b=rect(s,bx,by,3.35,1.02,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.2)
    t=txt(s,bx+0.12,by+0.07,3.1,0.4,ti,size=14,bold=True,color=col)
    d=txt(s,bx+0.12,by+0.5,3.1,0.46,de,size=12,color=DARK)
    anim.append([(ring,'appear'),(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 8 — Generiranje plana (podgraf)
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Kako nastaje plan putovanja?",
       "Plan ne prolazi odmah — iterativna validacija i popravak",color=ORANGE)

img(s,F+"nodes-2.png",0.28,1.5,w=12.75)

sub=[
    (ORANGE,LORANGE,"1. Generiranje",
     "2× LLM paralelno (claude-sonnet)\nkompaktnost vs. raznolikost"),
    (PURPLE,LPUR,"2. Provjera (bez LLM!)",
     "Kod ispituje preklapanja,\nudaljenosti, obvezna mjesta"),
    (PURPLE,LPUR,"3. Popravak",
     "LLM dobiva popis grešaka\ni popravlja samo te dijelove"),
    (GREEN,LGREEN,"4. Bodovanje",
     "6 kriterija: ruta 25%,\nprioritet 25%, kvaliteta 15%..."),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(sub):
    col_i=i%2; row_i=i//2
    x=0.28+col_i*6.55; y=5.5+row_i*0.95
    b=rect(s,x,y,6.4,0.88,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.0)
    t=txt(s,x+0.12,y+0.06,2.4,0.38,ti,size=14,bold=True,color=col)
    d=txt(s,x+2.6,y+0.06,3.7,0.75,de,size=13,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 9 — Inženjering upute i samoprovjera
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Kako smo poboljšali kvalitetu plana?",
       "Inženjering upute — kvaliteta plana ovisi o tome ŠTO kažemo modelu",color=INDIGO)

items=[
    (CORAL,LCORAL,"Problem: Preklapanja termina",
     "Svaki upit zahtijevao popravak\n→ plan kasni +34 sekunde"),
    (INDIGO,LIND,"Rješenje: Samoprovjera u uputi",
     'Model sam izračuna kraj svake\naktivnosti i pomakne preklapanja'),
    (ORANGE,LORANGE,"Problem: Nelogičan redoslijed",
     "Ruta: Opéra→Lafayette→Place Vendôme\n= vraćanje na sjever pa opet na jug"),
    (GREEN,LGREEN,"Rješenje: Smještaj kao polazište",
     '"Kreći od smještaja u jedan\ngeografski smjer kroz cijeli dan"'),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(items):
    col_i=i%2; row_i=i//2
    x=0.28+col_i*6.55; y=1.55+row_i*2.75
    b=rect(s,x,y,6.4,2.6,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.5)
    t=txt(s,x+0.18,y+0.15,6.0,0.5,ti,size=18,bold=True,color=col)
    d=txt(s,x+0.18,y+0.7,6.0,1.75,de,size=17,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])

res=lbl(s,0.28,7.02,12.75,0.42,
    "Rezultat: Plan prolazi iz prvog pokušaja → 295s umjesto 329s  |  Geografski smislena ruta",
    size=16,bold=False,color=WHITE,bg=INDIGO,rnd=True)
anim.append([(res,'fade')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 10 — Demo: Registracija & kreiranje
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Demo: Registracija i kreiranje putovanja",
       "Korisnik se registrira, potvrdi e-mail i odmah kreira putovanje",color=INDIGO)

img(s,F+"registracija.png",0.28,1.5,w=4.1)
img(s,F+"new1.png",4.5,1.5,w=4.1)
img(s,F+"new2.png",8.72,1.5,w=4.1)

for cx,cap,col,bg in [(2.0,"Registracija",INDIGO,LIND),
                       (6.5,"Kreiranje putovanja",INDIGO,LIND),
                       (10.8,"Prijevoz i slika",INDIGO,LIND)]:
    lbl(s,cx-1.55,7.08,3.1,0.38,cap,size=16,bold=True,color=col,bg=bg,rnd=True)

# ═══════════════════════════════════════════════════════════════
# SLIDE 11 — Demo: Mjesta i karta
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Demo: Zajednička karta i mjesta",
       "Dodaj mjesto — vidljivo svim sudionicima odmah (WebSocket)",color=TEAL)

img(s,F+"cijeli_trip.png",0.28,1.5,w=6.9)
img(s,F+"comments.png",7.28,1.5,w=5.85)

for cx,cap,col,bg in [(3.7,"Karta, mjesta i bočna traka",TEAL,LTEAL),
                       (10.2,"Komentari sudionika",PURPLE,LPUR)]:
    lbl(s,cx-2.6,7.08,5.2,0.38,cap,size=16,bold=True,color=col,bg=bg,rnd=True)

# ═══════════════════════════════════════════════════════════════
# SLIDE 12 — Demo: Agent chat
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Demo: Razgovor s agentom","Prirodni jezik → agentski sustav → strukturirani odgovor",color=PURPLE)

img(s,F+"chat.png",0.28,1.5,w=5.6)

caps=[
    (PURPLE,LPUR,"🤖  Preporuke mjesta",
     "Tavily web search + kontekst putovanja\n→ stvarne preporuke s ocjenama"),
    (TEAL,LTEAL,"✈️  Prijevoz",
     "Skyscanner za letove + web za vlakove\n→ stvarne cijene i polasci"),
    (ORANGE,LORANGE,"🏨  Smještaj",
     "Booking.com + Airbnb filtrirano\npo udaljenosti i ocjeni"),
    (GREEN,LGREEN,"📅  Plan putovanja",
     "Dnevni raspored s validacijom\nKlik 'Spremi' → dodaje u itinerar"),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(caps):
    y=1.5+i*1.48
    b=rect(s,6.15,y,7.0,1.35,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.0)
    t=txt(s,6.32,y+0.1,6.62,0.48,ti,size=18,bold=True,color=col)
    d=txt(s,6.32,y+0.62,6.62,0.65,de,size=16,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 13 — Demo: Dnevni raspored
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Demo: Dnevni raspored","Plan po danima — aktivnosti, prijevoz, obroci",color=GREEN)

img(s,F+"plan.png",4.38,1.5,w=4.6)

pts=[
    (GREEN,LGREEN,"Po danima","Navigacija strelicom — svaki dan zasebno"),
    (TEAL,LTEAL,"Vremena i trajanja","Početak, trajanje i put do sljedeće lokacije"),
    (INDIGO,LIND,"Karta uz plan","Markeri pokazuju redoslijed obilaska"),
    (PURPLE,LPUR,"Iterativna izmjena",'"Makni Ozio" → agent ažurira bez tog mjesta'),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(pts):
    y=1.5+i*1.48
    b=rect(s,0.28,y,3.95,1.35,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.0)
    t=txt(s,0.45,y+0.1,3.6,0.48,ti,size=18,bold=True,color=col)
    d=txt(s,0.45,y+0.62,3.6,0.65,de,size=16,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])

b2=rect(s,9.28,1.5,3.92,5.9,fill=LGRAY,rnd=True)
b2.line.color.rgb=RGBColor(0xCB,0xD5,0xE1); b2.line.width=Pt(0.5)
t2=txt(s,9.45,1.68,3.6,0.45,"Plan u chatu:",size=16,bold=True,color=GRAY)
plan_txt=("Vaša baza: Hôtel Beige – Paris 9e\n\n"
          "Dan 1  Lafayette · Opéra · Place Vendôme\n\n"
          "Dan 2  Louvre · Pyramides · Notre-Dame\n\n"
          "Dan 3  Champs-Élysées · Eiffelov toranj\n\n"
          "Dan 4  Invalides · Le Petit Cler\n\n"
          "Dan 5  Café Pouchkine · povratak")
t3=txt(s,9.45,2.18,3.6,5.0,plan_txt,size=14,color=DARK)
anim.append([(b2,'fade'),(t2,'appear'),(t3,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 14 — Demo: Dokumenti
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Demo: Dokumenti i karte",
       "Sve rezervacije i karte na jednom mjestu — dostupne svim sudionicima",color=AMBER)

img(s,F+"tickets.png",3.5,1.5,w=7.5)

cats=[
    (INDIGO,LIND,"🏨  Smještaj","Potvrde rezervacija hotela"),
    (TEAL,LTEAL,"🎭  Aktivnosti","Ulaznice, koncerti, ture"),
    (ORANGE,LORANGE,"✈️  Prijevoz","Boarding karte, vlak, bus"),
    (GRAY,LGRAY,"📄  Ostalo","Putovnica, osiguranje"),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(cats):
    y=1.5+i*1.48
    b=rect(s,0.28,y,3.08,1.35,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.0)
    t=txt(s,0.45,y+0.1,2.75,0.48,ti,size=17,bold=True,color=col)
    d=txt(s,0.45,y+0.62,2.75,0.65,de,size=15,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 15 — Evaluacija
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Eksperimentalna evaluacija",
       "Testni scenarij: Pariz, 5 dana, 18 mjesta, 2 sudionika",color=INDIGO)

hdrs=["Komponenta","Problem","Rješenje","Učinak"]
xs=[0.28,3.18,6.42,10.12]; ws=[2.82,3.12,3.58,3.05]
for j,(h,x,w) in enumerate(zip(hdrs,xs,ws)):
    lbl(s,x,1.45,w-0.06,0.55,h,size=16,bold=True,color=WHITE,bg=INDIGO)

rows=[
    ("Samoprovjera\nu uputi",
     "Plan pada na provjeri\n→ svaki upit zahtijeva\npopravak",
     "Model sam ispravlja\npreklapanja prije\npredaje na provjeru",
     "295s umjesto 329s\nProlaz iz 1. pokušaja ✓"),
    ("Obvezna\nmjesta",
     "Agent vraća mjesta\nkoja korisnik\neksplicitno traži izbrisati",
     "Čvor prepoznaje\nisključenje → makne iz\nliste obveznih",
     "Ponašanje usklađeno\ns namjerom\nkorisnika ✓"),
    ("Redoslijed\nobilaska",
     "Ruta se vraća\nunatrag (nepotrebno\nhodanje)",
     "Smještaj kao polazna\ntočka + kretanje\nu jednom smjeru",
     "Geografski smislena\nruta bez\npovrataka ✓"),
    ("Izbor\nmodela",
     "gemini-2.5-flash:\n13+ prekršaja,\nlošiji planovi",
     "claude-sonnet-4-6:\nvaljani planovi\niz 1. prolaza",
     "Sonnet: dobar omjer\ncijene i kvalitete;\nOpus: ~50s ✓"),
]
anim=[]
for i,row in enumerate(rows):
    y=2.08+i*1.28
    bg=LIND if i%2==0 else WHITE
    grp=[]
    for j,(cell,x,w) in enumerate(zip(row,xs,ws)):
        b=rect(s,x,y,w-0.06,1.2,fill=LPUR if j==0 else bg)
        b.line.color.rgb=LGRAY; b.line.width=Pt(0.5)
        tf=b.text_frame; tf.word_wrap=True
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.LEFT
        r=p.add_run(); r.text=cell; r.font.size=Pt(12)
        r.font.bold=(j==0); r.font.name=FONT
        r.font.color.rgb=PURPLE if j==0 else DARK
        grp.append((b,'fade'))
    anim.append(grp)
animate(s,anim)

# ═══════════════════════════════════════════════════════════════
# SLIDE 16 — Optimizacija rute
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
header(s,"Evaluacija: Optimizacija redoslijeda obilaska",
       "Primjer kako dodavanje smještaja ispravlja logiku rute",color=CORAL)

lbl(s,0.28,1.45,1.2,0.52,"❌ Prije",size=16,bold=True,color=CORAL,bg=LCORAL,rnd=True)
txt(s,1.6,1.5,5.0,0.45,"Bez smještaja — ruta se vraća unatrag",size=18,bold=True,color=DARK)
img(s,F+"lose.png",0.28,2.1,h=4.6)

lbl(s,7.05,1.45,1.2,0.52,"✅ Nakon",size=16,bold=True,color=GREEN,bg=LGREEN,rnd=True)
txt(s,8.38,1.5,5.0,0.45,"Sa smještajem — jednosmjerna ruta",size=18,bold=True,color=DARK)
img(s,F+"dobro.png",7.05,2.1,h=4.6)

lbl(s,0.28,6.95,12.75,0.5,
    'Rješenje: adresa smještaja u kontekst + uputa: "Kreći od smještaja, napreduj u jedan geografski smjer"',
    size=15,color=WHITE,bg=INDIGO,rnd=True)

# ═══════════════════════════════════════════════════════════════
# SLIDE 17 — Zaključak
# ═══════════════════════════════════════════════════════════════
s=sl()
rect(s,0,0,13.33,7.5,fill=WHITE)
rect(s,0,0,0.12,7.5,fill=INDIGO)
txt(s,0.3,0.15,12.6,0.8,"Zaključak",size=34,bold=True,color=DARK)
rect(s,0.3,0.95,12.75,0.04,fill=LGRAY)

zaklj=[
    (INDIGO,LIND,"🏗️  Implementiran cijeli sustav",
     "Rails 8 + LangGraph + PostgreSQL + vanjski API-ji\n"
     "— radi u stvarnom vremenu za više korisnika"),
    (PURPLE,LPUR,"🤖  Agentski pristup se pokazao boljim",
     "Samoprovjera + deterministička validacija + petlja popravka\n"
     "— planovi prolaze provjeru iz prvog pokušaja"),
    (GREEN,LGREEN,"📊  Evaluacija na stvarnom scenariju",
     "Pariz, 5 dana, 18 mjesta — svaki mehanizam testiran\n"
     "— Claude modeli daleko ispred Gemini-a"),
    (ORANGE,LORANGE,"🔮  Što dalje?",
     "Pouzdaniji API-ji za smještaj,\n"
     "podrška za više smještaja unutar putovanja"),
]
anim=[]
for i,(col,bg,ti,de) in enumerate(zaklj):
    ci=i%2; ri=i//2
    x=0.3+ci*6.52; y=1.15+ri*2.88
    b=rect(s,x,y,6.38,2.68,fill=bg,rnd=True)
    b.line.color.rgb=col; b.line.width=Pt(1.5)
    t=txt(s,x+0.2,y+0.15,5.95,0.55,ti,size=18,bold=True,color=col)
    d=txt(s,x+0.2,y+0.72,5.95,1.8,de,size=16,color=DARK)
    anim.append([(b,'fade'),(t,'appear'),(d,'appear')])
animate(s,anim)

lbl(s,0.3,7.05,12.72,0.42,
    "Hvala na pažnji!   •   Pitanja?   •   lorenajakic@fer.hr",
    size=20,bold=True,color=WHITE,bg=INDIGO,rnd=True)

# ── save ─────────────────────────────────────────────────────
out="/Users/lorenajakic/Desktop/diplomski_prezentacija.pptx"
prs.save(out)
print(f"✓ Saved {out}  ({len(prs.slides)} slajdova)")
